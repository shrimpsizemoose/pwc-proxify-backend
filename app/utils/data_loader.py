"""Data loading utilities - parse all data formats"""
import json
import csv
from pathlib import Path
from typing import Dict, List, Any
import pandas as pd
from email import policy
from email.parser import BytesParser
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


class DataLoader:
    """Load and parse all hackathon data sources"""

    def __init__(self, base_path: str = ".."):
        self.base_path = Path(base_path)
        self._cache: Dict[str, Any] = {}

    def load_salesforce_data(self) -> Dict[str, pd.DataFrame]:
        """Load all Salesforce CSV files"""
        if "salesforce" in self._cache:
            return self._cache["salesforce"]

        sf_path = self.base_path / "hackathon_fake_salesforce"
        data = {}

        try:
            data["accounts"] = pd.read_csv(sf_path / "Accounts.csv")
            data["contacts"] = pd.read_csv(sf_path / "Contacts.csv")
            data["opportunities"] = pd.read_csv(sf_path / "Opportunities.csv")
            data["activities"] = pd.read_csv(sf_path / "Activities.csv")

            # Load JSON files
            with open(sf_path / "ChatterPosts.json", "r") as f:
                data["chatter"] = json.load(f)

            with open(sf_path / "DataDictionary.json", "r") as f:
                data["data_dict"] = json.load(f)

        except Exception as e:
            print(f"Error loading Salesforce data: {e}")
            # Return empty dataframes
            data = {
                "accounts": pd.DataFrame(),
                "contacts": pd.DataFrame(),
                "opportunities": pd.DataFrame(),
                "activities": pd.DataFrame(),
                "chatter": [],
                "data_dict": {},
            }

        self._cache["salesforce"] = data
        return data

    def load_emails(self) -> List[Dict[str, Any]]:
        """Parse .eml email files"""
        if "emails" in self._cache:
            return self._cache["emails"]

        emails = []
        email_path = self.base_path / "hackathon_fake_mail_teams" / "mail"

        if not email_path.exists():
            return emails

        for eml_file in email_path.glob("*.eml"):
            try:
                with open(eml_file, "rb") as f:
                    msg = BytesParser(policy=policy.default).parse(f)

                emails.append({
                    "file": eml_file.name,
                    "from": str(msg.get("From", "")),
                    "to": str(msg.get("To", "")),
                    "subject": str(msg.get("Subject", "")),
                    "date": str(msg.get("Date", "")),
                    "body": msg.get_body(preferencelist=("plain")).get_content()
                    if msg.get_body(preferencelist=("plain"))
                    else "",
                })
            except Exception as e:
                print(f"Error parsing {eml_file}: {e}")

        self._cache["emails"] = emails
        return emails

    def load_teams_chats(self) -> List[Dict[str, Any]]:
        """Load Teams chat JSON files"""
        if "teams" in self._cache:
            return self._cache["teams"]

        chats = []
        teams_path = self.base_path / "hackathon_fake_mail_teams" / "teams"

        if not teams_path.exists():
            return chats

        for json_file in teams_path.glob("*.json"):
            try:
                with open(json_file, "r") as f:
                    chat_data = json.load(f)
                    chat_data["file"] = json_file.name
                    chats.append(chat_data)
            except Exception as e:
                print(f"Error loading {json_file}: {e}")

        self._cache["teams"] = chats
        return chats

    def load_sharepoint_docs(self) -> Dict[str, Any]:
        """Parse SharePoint documents (docx, xlsx, pptx)"""
        if "sharepoint" in self._cache:
            return self._cache["sharepoint"]

        docs = {"docx": [], "xlsx": [], "pptx": []}
        sp_path = self.base_path / "hackathon_fake_sharepoint_news" / "sharepoint"

        if not sp_path.exists():
            return docs

        # Parse DOCX files
        for docx_file in sp_path.glob("*.docx"):
            try:
                doc = Document(docx_file)
                text = "\n".join([para.text for para in doc.paragraphs])
                docs["docx"].append({
                    "file": docx_file.name,
                    "content": text,
                    "paragraphs": len(doc.paragraphs),
                })
            except Exception as e:
                # Fallback: try reading as plain text
                try:
                    with open(docx_file, "r", encoding="utf-8") as f:
                        text = f.read()
                    docs["docx"].append({
                        "file": docx_file.name,
                        "content": text,
                        "paragraphs": len(text.split("\n")),
                    })
                except:
                    print(f"Error parsing {docx_file}: {e}")

        # Parse XLSX files
        for xlsx_file in sp_path.glob("*.xlsx"):
            try:
                df = pd.read_excel(xlsx_file)
                docs["xlsx"].append({
                    "file": xlsx_file.name,
                    "data": df.to_dict(orient="records"),
                    "columns": list(df.columns),
                })
            except Exception as e:
                # Fallback: try reading as CSV
                try:
                    df = pd.read_csv(xlsx_file)
                    docs["xlsx"].append({
                        "file": xlsx_file.name,
                        "data": df.to_dict(orient="records"),
                        "columns": list(df.columns),
                    })
                except:
                    print(f"Error parsing {xlsx_file}: {e}")

        # Parse PPTX files
        for pptx_file in sp_path.glob("*.pptx"):
            try:
                prs = Presentation(pptx_file)
                slides_text = []
                for slide in prs.slides:
                    slide_text = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            slide_text.append(shape.text)
                    slides_text.append("\n".join(slide_text))

                docs["pptx"].append({
                    "file": pptx_file.name,
                    "slides": slides_text,
                    "slide_count": len(prs.slides),
                })
            except Exception as e:
                # Fallback: try reading as plain text
                try:
                    with open(pptx_file, "r", encoding="utf-8") as f:
                        text = f.read()
                    slides = [s.strip() for s in text.split("Slide ") if s.strip()]
                    docs["pptx"].append({
                        "file": pptx_file.name,
                        "slides": slides,
                        "slide_count": len(slides),
                    })
                except:
                    print(f"Error parsing {pptx_file}: {e}")

        self._cache["sharepoint"] = docs
        return docs

    def load_news_feed(self) -> List[Dict[str, Any]]:
        """Load news feed JSON"""
        if "news" in self._cache:
            return self._cache["news"]

        news_path = (
            self.base_path / "hackathon_fake_sharepoint_news" / "news_feed" / "NewsFeed.json"
        )

        try:
            with open(news_path, "r") as f:
                data = json.load(f)
                # Extract items array from JSON structure
                news = data.get("items", []) if isinstance(data, dict) else data
        except Exception as e:
            print(f"Error loading news feed: {e}")
            news = []

        self._cache["news"] = news
        return news

    def load_regulatory_feed(self) -> List[Dict[str, Any]]:
        """Load regulatory feed JSON"""
        if "regulatory" in self._cache:
            return self._cache["regulatory"]

        reg_path = (
            self.base_path
            / "hackathon_extra_calendar_knowledge_regulatory"
            / "regulatory_feed"
            / "RegulatoryFeed.json"
        )

        try:
            with open(reg_path, "r") as f:
                data = json.load(f)
                # Extract updates array from JSON structure
                regulatory = data.get("updates", []) if isinstance(data, dict) else data
        except Exception as e:
            print(f"Error loading regulatory feed: {e}")
            regulatory = []

        self._cache["regulatory"] = regulatory
        return regulatory

    def load_knowledge_hub(self) -> List[Dict[str, str]]:
        """Load knowledge hub markdown files"""
        if "knowledge" in self._cache:
            return self._cache["knowledge"]

        knowledge = []
        kb_path = (
            self.base_path
            / "hackathon_extra_calendar_knowledge_regulatory"
            / "knowledge_hub"
        )

        if not kb_path.exists():
            return knowledge

        for md_file in kb_path.glob("*.md"):
            try:
                with open(md_file, "r", encoding="utf-8") as f:
                    content = f.read()
                    knowledge.append({
                        "file": md_file.name,
                        "title": md_file.stem.replace("_", " "),
                        "content": content,
                    })
            except Exception as e:
                print(f"Error loading {md_file}: {e}")

        self._cache["knowledge"] = knowledge
        return knowledge

    def load_calendar(self) -> str:
        """Load calendar ICS file"""
        if "calendar" in self._cache:
            return self._cache["calendar"]

        cal_path = (
            self.base_path
            / "hackathon_extra_calendar_knowledge_regulatory"
            / "calendar"
            / "Meetings.ics"
        )

        try:
            with open(cal_path, "r", encoding="utf-8") as f:
                calendar = f.read()
        except Exception as e:
            print(f"Error loading calendar: {e}")
            calendar = ""

        self._cache["calendar"] = calendar
        return calendar

    def get_client_by_name(self, client_name: str) -> Dict[str, Any]:
        """Get all data for a specific client"""
        sf_data = self.load_salesforce_data()

        # Find account
        accounts = sf_data["accounts"]
        if accounts.empty:
            return None

        account = accounts[accounts["AccountName"].str.contains(client_name, case=False, na=False)]

        if account.empty:
            return None

        account_id = account.iloc[0]["AccountId"]

        # Get related contacts
        contacts = sf_data["contacts"]
        client_contacts = (
            contacts[contacts["AccountId"] == account_id].to_dict(orient="records")
            if not contacts.empty
            else []
        )

        # Get related opportunities
        opportunities = sf_data["opportunities"]
        client_opps = (
            opportunities[opportunities["AccountId"] == account_id].to_dict(orient="records")
            if not opportunities.empty
            else []
        )

        # Get related activities (join through opportunities)
        activities = sf_data["activities"]
        if not activities.empty and client_opps:
            # Get list of opportunity IDs for this client
            opp_ids = [opp["OpportunityId"] for opp in client_opps]
            # Filter activities by opportunity IDs
            client_activities = activities[activities["OpportunityId"].isin(opp_ids)].to_dict(orient="records")
        else:
            client_activities = []

        return {
            "account": account.to_dict(orient="records")[0],
            "contacts": client_contacts,
            "opportunities": client_opps,
            "activities": client_activities,
        }

    def clear_cache(self):
        """Clear data cache"""
        self._cache = {}
